"""T-1124-cto: Bundesvergleich aus KWRA 2021, TB 6, Kap. 7 (Checkliste Zeile 10, A2/A6/A7/A9).

Prüft ``app.data.kwra_kap7_risiko.BEREICHSVERGLEICH_RISIKO``:
(a) genau die fünf Systembereiche, in der Reihenfolge ``catalog.KWRA_SYSTEMBEREICHE``,
(b) jeder der vier Blöcke trägt Seiten, alle innerhalb von Kap. 7 (S. 146–154),
(c) Handlungserfordernisse 11/5, 10/6, 7/7, 0/2, 3/3 (S. 147–151),
(d) dieselben Zahlen ergeben sich aus der Arbeitsmappe (Spalten „Systembereich“ und
    „Handlungserfordernis“) — Gegenlesung, keine Quelle der Werte,
(e) Angaben zu hoch bewerteten Klimawirkungen: gültige Zeitscheibe und Fall, Wortlaut und
    Seite; die im Text genannten Zahlen stehen wörtlich da,
(f) Gewissheit und Einflüsse: wo der Text schweigt, steht None bzw. nichts — nie ein Wert,
(g) Zahl der Klimawirkungen je Bereich 30/31/23/7/9 (S. 146–151).
"""

from __future__ import annotations

import collections
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.data import catalog  # noqa: E402
from app.data.kwra_kap7_risiko import (  # noqa: E402
    ABBILDUNG_24,
    BEREICHSVERGLEICH_RISIKO,
    FAELLE,
    QUELLE,
    ZEITSCHEIBEN,
)

KAP7 = set(range(146, 155))
BLOECKE = ("risiko_ohne_anpassung", "gewissheit", "handlungserfordernisse",
           "klimatische_einfluesse")
MAPPE = os.path.join(
    os.path.dirname(__file__), "..", "..", "docs", "KWAR", "KWRA-2021_Klimawirkungen.xlsx"
)


def test_a_fuenf_bereiche_in_katalogreihenfolge():
    assert tuple(BEREICHSVERGLEICH_RISIKO) == catalog.KWRA_SYSTEMBEREICHE
    assert "Kapitel 7" in QUELLE and "S. 146–154" in QUELLE
    assert ABBILDUNG_24["seite"] == 152 and ABBILDUNG_24["seite"] in KAP7


def test_b_jeder_block_mit_seiten_aus_kap7():
    for bereich, v in BEREICHSVERGLEICH_RISIKO.items():
        for k in BLOECKE:
            seiten = v[k]["seiten"]
            assert seiten, (bereich, k)
            assert set(seiten) <= KAP7, (bereich, k, seiten)


def test_c_handlungserfordernisse_je_bereich():
    werte = [(BEREICHSVERGLEICH_RISIKO[b]["handlungserfordernisse"]["sehr_dringend"],
              BEREICHSVERGLEICH_RISIKO[b]["handlungserfordernisse"]["dringend"])
             for b in catalog.KWRA_SYSTEMBEREICHE]
    assert werte == [(11, 5), (10, 6), (7, 7), (0, 2), (3, 3)]


def test_d_handlungserfordernisse_gleich_arbeitsmappe():
    import openpyxl

    ws = openpyxl.load_workbook(MAPPE, read_only=True)["Klimawirkungen"]
    kopf = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    assert kopf[12] == "Systembereich" and kopf[32] == "Handlungserfordernis"
    zaehlung = collections.Counter(
        (r[12], r[32]) for r in ws.iter_rows(min_row=3, values_only=True)
        if isinstance(r[0], int))
    for bereich, v in BEREICHSVERGLEICH_RISIKO.items():
        he = v["handlungserfordernisse"]
        assert zaehlung[(bereich, "sehr dringend")] == he["sehr_dringend"], bereich
        assert zaehlung[(bereich, "dringend")] == he["dringend"], bereich


def test_e_hoch_bewertete_klimawirkungen_wie_im_text():
    for bereich, v in BEREICHSVERGLEICH_RISIKO.items():
        r = v["risiko_ohne_anpassung"]
        assert r["vergleich"], bereich
        assert r["hoch_bewertet"], bereich
        for a in r["hoch_bewertet"]:
            assert a["zeitscheibe"] in ZEITSCHEIBEN
            assert a["fall"] is None or a["fall"] in FAELLE
            assert a["wortlaut"] and a["seite"] in r["seiten"]
    kurz = {b: [(a["zeitscheibe"], a["anzahl_hoch"], a["anteil_prozent"])
                for a in v["risiko_ohne_anpassung"]["hoch_bewertet"]]
            for b, v in BEREICHSVERGLEICH_RISIKO.items()}
    assert kurz == {
        "Natürliche Systeme und Ressourcen": [("mitte", None, 60), ("ende", None, 70)],
        "Naturnutzende Wirtschaftssysteme": [("mitte", 9, None), ("ende", 16, None)],
        "Infrastrukturen und Gebäude": [("mitte", 6, None), ("ende", 8, None)],
        "Naturferne Wirtschaftssysteme": [("mitte", 1, None), ("ende", 2, None)],
        "Menschen und soziale Systeme": [("gegenwart", 1, None), ("mitte", None, None),
                                         ("ende", None, None)],
    }


def test_f_gewissheit_und_einfluesse_ohne_erfundene_werte():
    for bereich, v in BEREICHSVERGLEICH_RISIKO.items():
        g = v["gewissheit"]
        assert g["risiko_ohne_anpassung"] or g["anpassungskapazitaet"], bereich
        assert v["klimatische_einfluesse"]["einfluesse"], bereich
    # Zur Gewissheit der Anpassungskapazität sagt der Text nur bei zwei Bereichen etwas.
    mit_kapazitaet = [b for b, v in BEREICHSVERGLEICH_RISIKO.items()
                      if v["gewissheit"]["anpassungskapazitaet"] is not None]
    assert mit_kapazitaet == ["Natürliche Systeme und Ressourcen",
                              "Infrastrukturen und Gebäude"]


def test_g_zahl_der_klimawirkungen_je_bereich():
    assert [BEREICHSVERGLEICH_RISIKO[b]["anzahl_klimawirkungen"]
            for b in catalog.KWRA_SYSTEMBEREICHE] == [30, 31, 23, 7, 9]
