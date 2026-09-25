"""T-1123: KWRA-Systembereich je Klimawirkung (Checkliste Zeile 10, Anforderung A1).

Prüft ``app.data.kwra_systembereich_zuordnung.SYSTEMBEREICH_JE_KWRA_ID``:
(a) alle 102 KWRA-IDs sind zugeordnet, Werte nur aus den fünf Systembereichen,
(b) Zeile für Zeile gleich der Arbeitsmappe ``docs/KWAR/KWRA-2021_Klimawirkungen.xlsx``
    (Blatt ``Klimawirkungen``, Spalte „Systembereich“),
(c) jede Klimawirkung des Katalogs (aktiv und Roadmap) hat einen Systembereich, mit der
    Verteilung 16 / 15 / 13 / 2 / 6 über die fünf Bereiche,
(d) die Code-Zuordnung ``catalog.RISK_SYSTEMBEREICH`` der gerechneten Risiken stimmt mit
    dem Systembereich ihrer KWRA-ID überein,
(e) „Menschen und soziale Systeme“ umfasst neun Klimawirkungen (TB 6, S. 151).
"""

from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.data import catalog  # noqa: E402
from app.data.kwra_systembereich_zuordnung import (  # noqa: E402
    QUELLE,
    SYSTEMBEREICH_JE_KWRA_ID,
    systembereich_der_klimawirkung,
)

MAPPE = os.path.join(
    os.path.dirname(__file__), "..", "..", "docs", "KWAR", "KWRA-2021_Klimawirkungen.xlsx"
)


def _katalog_ids() -> set[int]:
    return ({p["kwra_id"] for p in catalog.PLANNED_RISKS}
            | {s["kwra_id"] for s in catalog.RISKS_BY_CODE.values()})


def test_a_alle_102_klimawirkungen_mit_zulaessigem_bereich():
    assert sorted(SYSTEMBEREICH_JE_KWRA_ID) == list(range(1, 103))
    assert set(SYSTEMBEREICH_JE_KWRA_ID.values()) == set(catalog.KWRA_SYSTEMBEREICHE)


def test_b_gleich_der_arbeitsmappe():
    import openpyxl

    ws = openpyxl.load_workbook(MAPPE, read_only=True)["Klimawirkungen"]
    kopf = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    assert kopf[0] == "ID" and kopf[12] == "Systembereich"
    mappe = {r[0]: r[12] for r in ws.iter_rows(min_row=3, values_only=True)
             if isinstance(r[0], int)}
    assert len(mappe) == 102
    abweichend = {i: (SYSTEMBEREICH_JE_KWRA_ID.get(i), v) for i, v in mappe.items()
                  if SYSTEMBEREICH_JE_KWRA_ID.get(i) != v}
    assert not abweichend, f"Abweichung zur Arbeitsmappe: {abweichend}"


def test_c_jede_klimawirkung_des_katalogs_hat_ihren_bereich():
    ids = _katalog_ids()
    assert len(ids) == 52
    fehlend = sorted(i for i in ids if i not in SYSTEMBEREICH_JE_KWRA_ID)
    assert not fehlend, f"Ohne Systembereich: {fehlend}"
    verteilung = [sum(systembereich_der_klimawirkung(i) == b for i in ids)
                  for b in catalog.KWRA_SYSTEMBEREICHE]
    assert verteilung == [16, 15, 13, 2, 6]


def test_d_code_zuordnung_passt_zur_kwra_id():
    for code, spec in catalog.RISKS_BY_CODE.items():
        assert catalog.RISK_SYSTEMBEREICH[code] == SYSTEMBEREICH_JE_KWRA_ID[spec["kwra_id"]], code


def test_e_menschen_und_soziale_systeme_neun_klimawirkungen():
    menschen = sorted(i for i, b in SYSTEMBEREICH_JE_KWRA_ID.items()
                      if b == "Menschen und soziale Systeme")
    assert menschen == [87, 95, 96, 97, 98, 99, 100, 101, 102]
    assert "KWRA-2021_Klimawirkungen.xlsx" in QUELLE
